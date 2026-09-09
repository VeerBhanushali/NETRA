"""Render docs/13-requirements-response.md as a Word document.

The response goes back to the people who sent the requirements as a
.docx, so it should arrive in the same format rather than as a markdown
file they have to render themselves.

Deliberately a converter for *this* document rather than a general
markdown-to-docx engine: it handles the constructs the response actually
uses (headings, tables, bullets, bold, inline code) and would need work
for anything else.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "13-requirements-response.md"
OUT = ROOT / "docs" / "NETRA-Requirements-Response.docx"

INK = RGBColor(0x1A, 0x1A, 0x1E)
SUBTLE = RGBColor(0x5E, 0x66, 0x72)
ACCENT = RGBColor(0x17, 0x50, 0x9C)
OK = RGBColor(0x0E, 0x6B, 0x38)
WARN = RGBColor(0x9A, 0x5B, 0x00)
DANGER = RGBColor(0xA8, 0x16, 0x16)


def shade(cell, hex_colour: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")          # never "solid" — renders black
    el.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def status_colour(text: str):
    t = text.lower()
    if "not built" in t:
        return DANGER
    if "partial" in t:
        return WARN
    if "built" in t:
        return OK
    return None


def add_runs(par, text: str, *, size=10.5, colour=INK) -> None:
    """Inline markdown: **bold**, `code`, and the arrow/×/≥ glyphs."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(size - 1)
            r.font.color.rgb = ACCENT
            continue
        else:
            r = par.add_run(part)
        r.font.size = Pt(size)
        r.font.color.rgb = colour


def main() -> int:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)

    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.8)

    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i]

        # --- tables -------------------------------------------------
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [[c.strip() for c in r.strip().strip("|").split("|")]
                    for r in block if not set(r) <= set("|-: ")]
            if not rows:
                continue
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for n, row in enumerate(rows):
                cells = table.add_row().cells
                for c, value in enumerate(row[:len(cells)]):
                    par = cells[c].paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    if n == 0:
                        add_runs(par, f"**{value.replace('**', '')}**", size=9.5)
                        shade(cells[c], "EEF2F7")
                    else:
                        add_runs(par, value, size=9,
                                 colour=status_colour(value) or INK)
            doc.add_paragraph()
            continue

        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        # --- headings -----------------------------------------------
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            if first_heading and level == 1:
                # Title block, not a body heading.
                t = doc.add_paragraph()
                t.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = t.add_run(text)
                r.bold = True
                r.font.size = Pt(19)
                r.font.color.rgb = ACCENT
                first_heading = False
            else:
                h = doc.add_heading(level=min(level, 3))
                for r in h.runs:
                    r.text = ""
                add_runs(h, text, size=14 - level, colour=ACCENT)
            i += 1
            continue

        # --- bullets ------------------------------------------------
        if stripped.startswith(("- ", "* ")):
            body = stripped[2:]
            while i + 1 < len(lines) and lines[i + 1].startswith("  ") \
                    and lines[i + 1].strip() and not lines[i + 1].strip().startswith(("-", "*", "|")):
                i += 1
                body += " " + lines[i].strip()
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, body)
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            body = re.sub(r"^\d+\.\s", "", stripped)
            while i + 1 < len(lines) and lines[i + 1].startswith("   ") and lines[i + 1].strip():
                i += 1
                body += " " + lines[i].strip()
            par = doc.add_paragraph(style="List Number")
            add_runs(par, body)
            i += 1
            continue

        # --- fenced code ---------------------------------------------
        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            for c in code:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.left_indent = Inches(0.25)
                r = par.add_run(c)
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                r.font.color.rgb = SUBTLE
            doc.add_paragraph()
            continue

        # --- body paragraph, joining wrapped lines --------------------
        body = stripped
        while i + 1 < len(lines) and lines[i + 1].strip() \
                and not lines[i + 1].startswith(("|", "#", "-", "*", "`")) \
                and not re.match(r"^\d+\.\s", lines[i + 1].strip()):
            i += 1
            body += " " + lines[i].strip()
        par = doc.add_paragraph()
        add_runs(par, body)
        i += 1

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}  ({OUT.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
